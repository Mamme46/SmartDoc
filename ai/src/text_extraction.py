from pathlib import Path

SUPPORTED_EXTENSIONS = ["txt", "md", "pdf", "docx"]


def extract_text(file_path) -> str:
    """
    Extrait le texte brut d'un fichier, quel que soit son format parmi
    SUPPORTED_EXTENSIONS. Lève ValueError si le format n'est pas supporté.
    """

    file_path = Path(file_path)
    extension = file_path.suffix.lower().lstrip(".")

    if extension in ("txt", "md"):
        return file_path.read_text(encoding="utf-8", errors="ignore")

    if extension == "pdf":
        return _extract_pdf(file_path)

    if extension == "docx":
        return _extract_docx(file_path)

    raise ValueError(f"Format non supporté : .{extension}")


def _extract_pdf(file_path: Path) -> str:
    from pypdf import PdfReader

    reader = PdfReader(str(file_path))
    pages = [page.extract_text() or "" for page in reader.pages]
    return "\n\n".join(pages)


def _extract_docx(file_path: Path) -> str:
    import docx

    document = docx.Document(str(file_path))
    paragraphs = [p.text for p in document.paragraphs]

    for table in document.tables:
        for row in table.rows:
            paragraphs.append(" | ".join(cell.text for cell in row.cells))

    return "\n".join(paragraphs)
